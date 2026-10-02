import io
import pytest
from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas
import pypdf

# Helper function to generate PDF bytes programmatically
def create_pdf_bytes(text_lines=None, encrypted=False, password="secret") -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    if text_lines is not None:
        y = 750
        for line in text_lines:
            c.drawString(50, y, line)
            y -= 20
            if y < 50:
                c.showPage()
                y = 750
    else:
        # Draw a shape instead of text (image-only / scanned mock)
        c.rect(50, 50, 500, 700)
    c.save()
    pdf_data = buffer.getvalue()

    if encrypted:
        reader = pypdf.PdfReader(io.BytesIO(pdf_data))
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.encrypt(password)
        encrypted_buffer = io.BytesIO()
        writer.write(encrypted_buffer)
        return encrypted_buffer.getvalue()

    return pdf_data


# Helper function to generate DOCX bytes programmatically
def create_docx_bytes(paragraphs=None, table_data=None) -> bytes:
    doc = Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def valid_resume_text_lines():
    return [
        "Jane Doe",
        "Email: jane.doe@example.com | Phone: +91 9876543210",
        "LinkedIn: linkedin.com/in/janedoe | GitHub: github.com/janedoe",
        "PROFESSIONAL SUMMARY:",
        "Accomplished Software Engineer with over 5 years of experience building scalable backend APIs,",
        "microservices, and data processing pipelines using Python, Flask, SQL, and Docker.",
        "WORK EXPERIENCE:",
        "Senior Backend Developer - Tech Corp (2021 - Present)",
        "Designed and implemented RESTful microservices processing 1M daily requests.",
        "Reduced database query latency by 35% through query optimization and caching.",
        "EDUCATION:",
        "Bachelor of Science in Computer Science - University of Technology (2017 - 2021)",
        "TECHNICAL SKILLS:",
        "Python, Flask, Django, PostgreSQL, Docker, Kubernetes, Git, REST API, Pytest",
        "PROJECTS:",
        "Smart Resume Analyzer - AI based resume parsing and scoring platform.",
        "CERTIFICATIONS:",
        "AWS Certified Solutions Architect Associate"
    ]


def test_upload_no_file(client):
    """Test 400 when no file part is provided."""
    response = client.post("/api/upload")
    assert response.status_code == 400
    data = response.get_json()
    assert data["error"]["code"] == "NO_FILE_PROVIDED"


def test_upload_empty_filename(client):
    """Test 400 when filename is empty."""
    data = {"file": (io.BytesIO(b"%PDF-1.4 test"), "")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    res_json = response.get_json()
    assert res_json["error"]["code"] == "NO_FILE_SELECTED"


def test_upload_zero_byte_file(client):
    """Test 400 when file is 0 bytes."""
    data = {"file": (io.BytesIO(b""), "empty.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    res_json = response.get_json()
    assert res_json["error"]["code"] == "ZERO_BYTE_FILE"


def test_upload_invalid_extension(client):
    """Test 400 when extension is not .pdf or .docx."""
    data = {"file": (io.BytesIO(b"some text content"), "resume.txt")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    res_json = response.get_json()
    assert res_json["error"]["code"] == "INVALID_EXTENSION"


def test_upload_mismatched_magic_bytes_pdf(client):
    """Test 400 when file has .pdf extension but invalid magic bytes."""
    fake_pdf = b"Plain text masquerading as PDF"
    data = {"file": (io.BytesIO(fake_pdf), "fake.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    res_json = response.get_json()
    assert res_json["error"]["code"] == "INVALID_MIME_TYPE"


def test_upload_mismatched_magic_bytes_docx(client):
    """Test 400 when file has .docx extension but invalid magic bytes."""
    fake_docx = b"Plain text masquerading as DOCX"
    data = {"file": (io.BytesIO(fake_docx), "fake.docx")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 400
    res_json = response.get_json()
    assert res_json["error"]["code"] == "INVALID_MIME_TYPE"


def test_upload_oversized_file(client, app):
    """Test 413 when file exceeds MAX_UPLOAD_MB."""
    # Temporarily set MAX_UPLOAD_MB to 1MB for test
    app.config["MAX_UPLOAD_MB"] = 1
    oversized_data = b"%PDF-1.4 " + (b"X" * (1024 * 1024 * 2))
    data = {"file": (io.BytesIO(oversized_data), "large.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 413
    res_json = response.get_json()
    assert res_json["error"]["code"] == "FILE_TOO_LARGE"


def test_upload_encrypted_pdf(client):
    """Test 422 when PDF is password protected."""
    encrypted_pdf = create_pdf_bytes(["Protected text content"], encrypted=True)
    data = {"file": (io.BytesIO(encrypted_pdf), "protected.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 422
    res_json = response.get_json()
    assert res_json["error"]["code"] == "ENCRYPTED_FILE"


def test_upload_scanned_image_only_pdf(client):
    """Test 422 when PDF has no selectable text."""
    scanned_pdf = create_pdf_bytes(text_lines=None) # empty canvas
    data = {"file": (io.BytesIO(scanned_pdf), "scanned.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 422
    res_json = response.get_json()
    assert res_json["error"]["code"] == "NO_TEXT_FOUND"


def test_upload_text_too_short(client):
    """Test 422 when extracted resume text is below MIN_WORD_COUNT."""
    short_pdf = create_pdf_bytes(["John Doe", "Short resume text"])
    data = {"file": (io.BytesIO(short_pdf), "short.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 422
    res_json = response.get_json()
    assert res_json["error"]["code"] == "TEXT_TOO_SHORT"


def test_upload_valid_pdf_resume(client, valid_resume_text_lines):
    """Test successful upload and parsing of a valid PDF resume."""
    pdf_bytes = create_pdf_bytes(valid_resume_text_lines)
    data = {"file": (io.BytesIO(pdf_bytes), "jane_doe_resume.pdf")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 201
    res_json = response.get_json()

    assert "resume_id" in res_json
    assert res_json["filename"] == "jane_doe_resume.pdf"
    assert res_json["word_count"] > 30
    assert res_json["page_count"] >= 1

    # Check extracted contacts
    contacts = res_json["contacts"]
    assert "jane.doe@example.com" in contacts["emails"]
    assert len(contacts["phones"]) > 0
    assert any("linkedin.com/in/janedoe" in link for link in contacts["linkedin"])
    assert any("github.com/janedoe" in gh for gh in contacts["github"])

    # Check detected sections
    sections = res_json["detected_sections"]
    assert "experience" in sections
    assert "education" in sections
    assert "skills" in sections


def test_upload_valid_docx_resume_with_table(client):
    """Test successful upload and parsing of a valid DOCX resume with paragraphs and tables."""
    paragraphs = [
        "Alex Mercer",
        "alex.mercer@devmail.com | +91 9988776655",
        "github.com/alexmercer | linkedin.com/in/alexmercer",
        "PROFESSIONAL SUMMARY:",
        "Full stack software developer with expertise in Web applications, REST APIs, microservices and automated testing.",
        "WORK EXPERIENCE:",
        "Full Stack Developer at InnovateTech Solutions (2020 - Present)",
        "Developed web applications using Python Flask framework and vanilla JavaScript.",
        "EDUCATION:",
        "B.Tech in Information Technology from State University (2016 - 2020)",
        "TECHNICAL SKILLS:"
    ]
    table_data = [
        ["Category", "Skills"],
        ["Languages", "Python, JavaScript, SQL, HTML, CSS"],
        ["Frameworks", "Flask, FastAPI, Pytest, Tailwind"],
        ["Tools", "Git, Docker, VSCode, Linux"]
    ]

    docx_bytes = create_docx_bytes(paragraphs=paragraphs, table_data=table_data)
    data = {"file": (io.BytesIO(docx_bytes), "alex_mercer.docx")}
    response = client.post("/api/upload", data=data, content_type="multipart/form-data")
    assert response.status_code == 201
    res_json = response.get_json()

    assert res_json["filename"] == "alex_mercer.docx"
    assert res_json["word_count"] > 30
    assert "alex.mercer@devmail.com" in res_json["contacts"]["emails"]

    # Table content should be included in raw_text and detected_sections
    raw_text = res_json["raw_text"]
    assert "FastAPI" in raw_text
    assert "Docker" in raw_text


def test_upload_hash_deduplication(client, valid_resume_text_lines):
    """Test that uploading the exact same file twice reuses the stored resume record."""
    pdf_bytes = create_pdf_bytes(valid_resume_text_lines)

    # First upload
    data1 = {"file": (io.BytesIO(pdf_bytes), "original_resume.pdf")}
    res1 = client.post("/api/upload", data=data1, content_type="multipart/form-data")
    assert res1.status_code == 201
    json1 = res1.get_json()
    resume_id1 = json1["resume_id"]

    # Second upload (same bytes)
    data2 = {"file": (io.BytesIO(pdf_bytes), "duplicate_resume.pdf")}
    res2 = client.post("/api/upload", data=data2, content_type="multipart/form-data")
    assert res2.status_code == 200
    json2 = res2.get_json()
    resume_id2 = json2["resume_id"]

    assert resume_id1 == resume_id2
    assert json1["file_hash"] == json2["file_hash"]
