import hashlib
import io
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any

import docx
import pdfplumber
import pypdf


# Custom Exceptions for Resume Upload & Parsing
class ResumeParsingError(Exception):
    """Base class for resume parsing errors with HTTP status code and error code."""
    def __init__(self, code: str, message: str, status_code: int = 400):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code

class InvalidExtensionError(ResumeParsingError):
    def __init__(self, message: str = "Unsupported file extension. Only .pdf and .docx files are allowed."):
        super().__init__(code="INVALID_EXTENSION", message=message, status_code=400)

class ZeroByteFileError(ResumeParsingError):
    def __init__(self, message: str = "Uploaded file is empty (0 bytes)."):
        super().__init__(code="ZERO_BYTE_FILE", message=message, status_code=400)

class MismatchedFileTypeError(ResumeParsingError):
    def __init__(self, message: str = "File extension does not match the actual file content."):
        super().__init__(code="INVALID_MIME_TYPE", message=message, status_code=400)

class FileTooLargeError(ResumeParsingError):
    def __init__(self, message: str = "File size exceeds maximum allowed upload limit."):
        super().__init__(code="FILE_TOO_LARGE", message=message, status_code=413)

class EncryptedFileError(ResumeParsingError):
    def __init__(self, message: str = "File is password protected or encrypted. Please upload an unencrypted file."):
        super().__init__(code="ENCRYPTED_FILE", message=message, status_code=422)

class UnreadableFileError(ResumeParsingError):
    def __init__(self, message: str = "Unable to read or parse file content. File may be corrupted."):
        super().__init__(code="UNREADABLE_FILE", message=message, status_code=422)

class NoTextFoundError(ResumeParsingError):
    def __init__(self, message: str = "No selectable text found in PDF. Scanned or image-only PDFs cannot be processed by ATS systems."):
        super().__init__(code="NO_TEXT_FOUND", message=message, status_code=422)

class TextTooShortError(ResumeParsingError):
    def __init__(self, message: str = "Extracted text is too short to be a valid resume."):
        super().__init__(code="TEXT_TOO_SHORT", message=message, status_code=422)


def compute_file_hash(file_bytes: bytes) -> str:
    """Computes SHA-256 hash of raw file bytes."""
    return hashlib.sha256(file_bytes).hexdigest()


def validate_file_content(filename: str, raw_bytes: bytes, max_upload_mb: int = 5) -> str:
    """Validates file extension, length, and magic bytes. Returns file extension (pdf/docx)."""
    if not filename or "." not in filename:
        raise InvalidExtensionError("No valid file extension found.")

    ext = filename.rsplit(".", 1)[1].lower()
    if ext not in {"pdf", "docx"}:
        raise InvalidExtensionError(f"Unsupported extension '.{ext}'. Only .pdf and .docx are allowed.")

    file_size = len(raw_bytes)
    if file_size == 0:
        raise ZeroByteFileError()

    max_bytes = max_upload_mb * 1024 * 1024
    if file_size > max_bytes:
        raise FileTooLargeError(f"File size ({file_size / (1024*1024):.1f}MB) exceeds limit of {max_upload_mb}MB.")

    # Validate Magic Bytes
    if ext == "pdf":
        if not raw_bytes.startswith(b"%PDF"):
            raise MismatchedFileTypeError("File extension is .pdf but magic bytes do not match PDF format.")
    elif ext == "docx":
        # PK Zip header for docx
        if not raw_bytes.startswith(b"PK\x03\x04"):
            raise MismatchedFileTypeError("File extension is .docx but magic bytes do not match DOCX package format.")

    return ext


def normalize_text(text: str) -> str:
    """Normalizes raw extracted text (Unicode normalization, line breaks, control chars, whitespace)."""
    if not text:
        return ""

    # Unicode normalization (NFKC)
    text = unicodedata.normalize("NFKC", text)

    # Strip control chars except \n and \t
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)

    # Fix hyphenated line breaks (e.g. devel-\nopment -> development)
    text = re.sub(r'(\b[a-zA-Z]{2,})-\s*\n\s*([a-zA-Z]{2,}\b)', r'\1\2', text)

    # Replace bullet glyphs with standard bullet
    bullet_pattern = r'[•▪►●➢❖–—*]\s*'
    text = re.sub(bullet_pattern, '• ', text)

    # Clean multiple spaces per line
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines()]

    # Collapse consecutive blank lines (max 2 newlines)
    normalized = "\n".join(lines)
    normalized = re.sub(r'\n{3,}', '\n\n', normalized)

    return normalized.strip()


def extract_contacts(text: str) -> dict[str, list[str]]:
    """Extracts email, phone, LinkedIn, and GitHub contacts via regex."""
    contacts = {
        "emails": [],
        "phones": [],
        "linkedin": [],
        "github": []
    }
    if not text:
        return contacts

    # Email extraction
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    emails = list(dict.fromkeys(re.findall(email_pattern, text)))
    contacts["emails"] = emails

    # Phone extraction (support international & Indian formats)
    phone_pattern = r'(?:\+?\d{1,3}[\s.-]?)?\(?\d{3,5}\)?[\s.-]?\d{3,5}[\s.-]?\d{4,5}'
    raw_phones = re.findall(phone_pattern, text)
    valid_phones = []
    for p in raw_phones:
        cleaned_p = p.strip()
        digits = re.sub(r'\D', '', cleaned_p)
        if 10 <= len(digits) <= 15:
            if cleaned_p not in valid_phones:
                valid_phones.append(cleaned_p)
    contacts["phones"] = valid_phones

    # LinkedIn extraction
    linkedin_pattern = r'(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+/?'
    linkedin_matches = re.findall(linkedin_pattern, text, re.IGNORECASE)
    cleaned_linkedin = []
    for link in linkedin_matches:
        url = link if link.startswith("http") else "https://" + link
        if url not in cleaned_linkedin:
            cleaned_linkedin.append(url)
    contacts["linkedin"] = cleaned_linkedin

    # GitHub extraction
    github_pattern = r'(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+/?'
    github_matches = re.findall(github_pattern, text, re.IGNORECASE)
    cleaned_github = []
    for gh in github_matches:
        url = gh if gh.startswith("http") else "https://" + gh
        if url not in cleaned_github:
            cleaned_github.append(url)
    contacts["github"] = cleaned_github

    return contacts


def load_section_synonyms(config_path: Path) -> dict[str, list[str]]:
    """Loads section synonym map from JSON config file."""
    if not config_path.exists():
        # Fallback default synonyms if file doesn't exist
        return {
            "contact": ["contact", "personal details", "contact info"],
            "summary": ["summary", "professional summary", "profile", "about me", "objective"],
            "experience": ["experience", "work experience", "employment history", "work history"],
            "education": ["education", "academic background", "educational qualifications"],
            "skills": ["skills", "technical skills", "core competencies", "key skills", "technologies"],
            "projects": ["projects", "key projects", "academic projects"],
            "certifications": ["certifications", "certificates", "licenses & certifications"]
        }
    with open(config_path, encoding="utf-8") as f:
        data = json.load(f)
        return data.get("synonyms", {})


def detect_sections(text: str, synonym_map: dict[str, list[str]]) -> dict[str, str]:
    """Detects resume sections based on synonym map matching."""
    sections = {
        "contact": "",
        "summary": "",
        "experience": "",
        "education": "",
        "skills": "",
        "projects": "",
        "certifications": "",
        "other": ""
    }

    if not text:
        return sections

    lines = text.splitlines()
    current_section = "contact"
    section_lines: dict[str, list[str]] = {key: [] for key in sections}

    # Reverse mapping for fast matching
    synonym_lookup = {}
    for sec_key, synonyms in synonym_map.items():
        for syn in synonyms:
            synonym_lookup[syn.lower().strip()] = sec_key

    for line in lines:
        clean_line = line.strip()
        if not clean_line:
            continue

        # Strip formatting marks (bullets, colons, markdown #, etc.) for header check
        header_candidate = re.sub(r'^[\#\*\-•\s]+|[\:\s]+$', '', clean_line).strip().lower()

        matched_section = None
        if header_candidate in synonym_lookup:
            matched_section = synonym_lookup[header_candidate]

        if matched_section:
            current_section = matched_section
        else:
            section_lines[current_section].append(clean_line)

    for sec_key in sections:
        sections[sec_key] = "\n".join(section_lines[sec_key]).strip()

    return sections


def parse_pdf(raw_bytes: bytes) -> tuple[str, int, list[str]]:
    """Parses PDF bytes using pdfplumber with pypdf fallback."""
    text_content = []
    page_count = 0
    warnings = []
    layout_risky = False

    # First attempt with pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                # Check for multi-column layout risk via word positioning
                words = page.extract_words()
                if len(words) > 10:
                    x0_vals = [w["x0"] for w in words]
                    min_x, max_x = min(x0_vals), max(x0_vals)
                    page_width = page.width or 600
                    # If text spans across distinct left and right columns
                    left_col = [w for w in words if w["x0"] < page_width * 0.45]
                    right_col = [w for w in words if w["x0"] > page_width * 0.55]
                    if len(left_col) > 15 and len(right_col) > 15:
                        layout_risky = True

                extracted = page.extract_text()
                if extracted:
                    text_content.append(extracted)

    except pdfplumber.pdfminer.pdfdocument.PDFPasswordIncorrect:
        raise EncryptedFileError()
    except Exception:
        # Fallback to pypdf
        try:
            reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            if reader.is_encrypted:
                raise EncryptedFileError()
            page_count = len(reader.pages)
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    text_content.append(txt)
        except EncryptedFileError:
            raise
        except pypdf.errors.FileNotDecryptedError:
            raise EncryptedFileError()
        except Exception as fallback_err:
            raise UnreadableFileError(f"Failed to parse PDF document: {fallback_err!s}")

    full_raw_text = "\n".join(text_content).strip()

    # If pdfplumber returned empty text, try pypdf fallback before declaring empty
    if not full_raw_text and page_count > 0:
        try:
            reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            page_count = len(reader.pages)
            fallback_text = []
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    fallback_text.append(txt)
            full_raw_text = "\n".join(fallback_text).strip()
        except Exception:
            pass

    if layout_risky:
        warnings.append("Multi-column layout detected. Side-by-side columns may affect ATS parser ordering.")

    return full_raw_text, page_count, warnings


def parse_docx(raw_bytes: bytes) -> tuple[str, int, list[str]]:
    """Parses DOCX bytes including paragraphs, headers, and tables."""
    text_content = []
    warnings = []

    try:
        doc = docx.Document(io.BytesIO(raw_bytes))
    except Exception as e:
        raise UnreadableFileError(f"Failed to open DOCX package: {e!s}")

    # Headers
    for section in doc.sections:
        if section.header and not section.header.is_linked_to_previous:
            for p in section.header.paragraphs:
                if p.text.strip():
                    text_content.append(p.text.strip())

    # Paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            text_content.append(p.text.strip())

    # Tables
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_str = cell.text.strip()
                if cell_str and cell_str not in row_text:
                    row_text.append(cell_str)
            if row_text:
                text_content.append(" | ".join(row_text))

    full_text = "\n".join(text_content).strip()
    words = full_text.split()
    word_count = len(words)
    estimated_pages = max(1, math.ceil(word_count / 400))

    return full_text, estimated_pages, warnings


def parse_resume_bytes(
    file_bytes: bytes,
    filename: str,
    sections_config_file: Path,
    max_upload_mb: int = 5,
    min_word_count: int = 30
) -> dict[str, Any]:
    """
    Main entrypoint to validate and parse resume bytes into structured dictionary.
    """
    # 1. Validate file format & magic bytes
    ext = validate_file_content(filename, file_bytes, max_upload_mb=max_upload_mb)

    # 2. Extract raw text & metadata
    if ext == "pdf":
        raw_text, page_count, warnings = parse_pdf(file_bytes)
    else:
        raw_text, page_count, warnings = parse_docx(file_bytes)

    # 3. Check for empty / scanned PDF text
    clean_text = raw_text.strip()
    if len(clean_text) < 15:
        raise NoTextFoundError("No selectable text found in the file. Scanned or image-only files cannot be processed by ATS systems.")

    # 4. Normalize text
    normalized_text = normalize_text(raw_text)

    # 5. Check word count threshold
    words = normalized_text.split()
    word_count = len(words)
    if word_count < min_word_count:
        raise TextTooShortError(
            f"Extracted text is too short to be a valid resume ({word_count} words found, minimum {min_word_count} required)."
        )

    # 6. Extract Contacts
    contacts = extract_contacts(normalized_text)

    # 7. Detect Sections
    synonyms = load_section_synonyms(sections_config_file)
    detected_sections = detect_sections(normalized_text, synonyms)

    return {
        "raw_text": normalized_text,
        "page_count": page_count,
        "word_count": word_count,
        "detected_sections": detected_sections,
        "contacts": contacts,
        "warnings": warnings
    }
